"""Run from repository root after installing backend dependencies."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1] / "sample_data"
ROOT.mkdir(exist_ok=True)
headers = [
    "transaction_id",
    "date",
    "store",
    "sku",
    "quantity",
    "unit_price",
    "discount",
]
valid = [
    [
        f"UPLOAD-OCT-{i:03}",
        "2026-10-18",
        f"ST{i % 3 + 1:02}",
        f"RP-{i:03}-1",
        3,
        1299,
        0.05,
    ]
    for i in range(1, 16)
]
invalid = [
    ["UPLOAD-BAD-001", "2026-10-18", "ST01", "INVALID-SKU", 2, 999, 0],
    ["UPLOAD-BAD-002", "", "ST01", "RP-001-1", 2, 999, 0],
    valid[0],
    ["UPLOAD-BAD-003", "2026-10-18", "ST01", "RP-001-1", "bad", 999, 0],
]


def style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="216C56")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 24


wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(headers)
for r in valid:
    ws.append(r)
style(ws)
ws = wb.create_sheet("Validation examples")
ws.append(headers)
for r in invalid:
    ws.append(r)
style(ws)
wb.save(ROOT / "sales_import.xlsx")
with (ROOT / "sales_template.csv").open("w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(headers)
    writer.writerows(valid[:2])
wb = Workbook()
ws = wb.active
ws.title = "Inventory"
ws.append(["date", "store", "sku", "quantity", "reserved", "received_date"])
for i in range(1, 11):
    ws.append(["2026-10-18", "ST01", f"RP-{i:03}-1", 5, 1, "2026-10-01"])
style(ws)
wb.save(ROOT / "inventory_import.xlsx")
print("Generated 3 fictional sample files. Sales XLSX: 15 valid, 4 deliberately rejected rows.")
