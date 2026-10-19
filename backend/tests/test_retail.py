import io
from datetime import timedelta

import numpy as np
import pytest
from openpyxl import load_workbook

from app.analytics import Filters, analytics, inventory_metrics
from app.core import SessionLocal
from app.forecasting import evaluate_series, predict
from app.imports import number, read_file
from app.models import Organization, PurchaseOrder, Workflow
from app.reports import safe_cell
from app.seed import AS_OF
from app.workflows import execute, recommendations, transition


def upload(client, headers, content, filename="sales.csv"):
    return client.post(
        "/api/v1/imports/preview",
        headers=headers,
        files={"file": (filename, content, "text/csv")},
    )


HEADER = "transaction_id,date,store,sku,quantity,unit_price,discount\n"
VALID = "NEW-1,2026-10-18,ST01,RP-001-1,3,100,0.1\n"


def test_health_and_auth(client):
    assert client.get("/api/v1/health").json()["status"] == "ok"
    assert client.get("/api/v1/analytics").status_code == 401


def test_known_kpis_and_returns(client, fixture_data):
    response = client.get("/api/v1/analytics", headers=fixture_data["analyst"])
    assert response.status_code == 200, response.text
    m = response.json()["metrics"]
    assert m["revenue"] == 5310  # 30 * 180 - 90
    assert m["gross_profit"] == 2950
    assert m["units"] == 59
    assert m["stock"] == 105
    assert m["inventory_value"] == 5200
    assert m["dead_stock_value"] == 5000
    assert m["low_stock"] == 1


def test_stock_cover_and_reorder(client, fixture_data):
    with SessionLocal() as db:
        cfg = db.get(Organization, fixture_data["org"]).settings
        rows = inventory_metrics(db, fixture_data["org"], Filters(), cfg)
        active = next(r for r in rows if r["sku"] == "RP-001-1")
        dead = next(r for r in rows if r["sku"] == "RP-002-1")
        assert active["daily_demand"] == 2
        assert active["cover_days"] == 2
        assert active["reorder_point"] == 34
        assert active["reorder_quantity"] == 30
        assert dead["cover_days"] is None
        assert dead["dead_stock"] is True
        assert dead["reorder_quantity"] == 0
        assert recommendations(db, fixture_data["org"], cfg) == 1
        assert recommendations(db, fixture_data["org"], cfg) == 0


def test_upload_commit_replay_and_updated_metrics(client, fixture_data):
    h = fixture_data["analyst"]
    preview = upload(client, h, (HEADER + VALID).encode())
    assert preview.status_code == 200, preview.text
    batch = preview.json()
    assert (batch["accepted"], batch["rejected"]) == (1, 0)
    path = f"/api/v1/imports/{batch['id']}/commit"
    assert client.post(path, headers=h).status_code == 200
    assert client.post(path, headers=h).json()["idempotent_replay"] is True
    assert upload(client, h, (HEADER + VALID).encode()).json()["id"] == batch["id"]
    assert client.get("/api/v1/analytics", headers=h).json()["metrics"]["revenue"] == 5580
    assert client.get("/api/v1/workflows", headers=h).json()[0]["status"] == "Completed"


@pytest.mark.parametrize(
    "bad",
    [
        "BAD,2026-10-18,ST01,INVALID,2,100,0",
        "BAD,,ST01,RP-001-1,2,100,0",
        "BAD,2026-10-18,ST01,RP-001-1,1.5,100,0",
        "BAD,2026-10-18,ST01,RP-001-1,2,-1,0",
        "BAD,2026-10-18,ST01,RP-001-1,2,100,1.2",
        "BAD,2026-10-18,ST01,RP-001-1,NaN,100,0",
        "FIX-0,2026-10-18,ST01,RP-001-1,2,100,0",
        "BAD,2099-01-01,ST01,RP-001-1,2,100,0",
    ],
)
def test_invalid_rows(client, fixture_data, bad):
    batch = upload(client, fixture_data["analyst"], (HEADER + bad + "\n").encode()).json()
    assert batch["accepted"] == 0
    assert batch["rejected"] == 1
    assert batch["errors"][0]["row"] == 2


def test_missing_headers_and_mapping(client, fixture_data):
    h = fixture_data["analyst"]
    result = upload(
        client,
        h,
        b"ref,when,location,item,count,price\nMAPPED,18/10/2026,ST01,RP-001-1,2,100\n",
    ).json()
    assert result["rejected"] == 1
    mapping = dict(
        zip(
            ["ref", "when", "location", "item", "count", "price"],
            ["transaction_id", "date", "store", "sku", "quantity", "unit_price"],
        )
    )
    response = client.post(
        f"/api/v1/imports/{result['id']}/validate", headers=h, json={"mapping": mapping}
    )
    assert response.json()["accepted"] == 1


def test_duplicates_within_file(client, fixture_data):
    result = upload(client, fixture_data["analyst"], (HEADER + VALID + VALID).encode()).json()
    assert result["accepted"] == 1 and result["rejected"] == 1


def test_secure_file_parsing(client, fixture_data):
    assert upload(client, fixture_data["analyst"], b"<script/>", "bad.html").status_code == 415
    assert upload(client, fixture_data["analyst"], b"not an xlsx", "bad.xlsx").status_code == 422
    with pytest.raises(Exception):
        read_file(b"x" * (8 * 1024 * 1024 + 1), "big.csv")
    with pytest.raises(ValueError):
        number("Infinity")


def test_tenant_isolation_and_approval_permissions(client, fixture_data):
    with SessionLocal() as db:
        recommendations(db, fixture_data["org"], db.get(Organization, fixture_data["org"]).settings)
        db.commit()
    rec = client.get("/api/v1/recommendations", headers=fixture_data["analyst"]).json()[0]
    assert client.get("/api/v1/recommendations", headers=fixture_data["outsider"]).json() == []
    assert (
        client.post(
            f"/api/v1/recommendations/{rec['id']}/submit",
            headers=fixture_data["outsider"],
        ).status_code
        == 404
    )
    approval = client.post(
        f"/api/v1/recommendations/{rec['id']}/submit", headers=fixture_data["analyst"]
    ).json()
    path = f"/api/v1/approvals/{approval['id']}/decision"
    body = {"decision": "approve", "reason": "Stock cover reviewed"}
    assert client.post(path, headers=fixture_data["analyst"], json=body).status_code == 403
    assert client.post(path, headers=fixture_data["outsider"], json=body).status_code == 404
    assert client.post(path, headers=fixture_data["manager"], json=body).status_code == 200
    assert client.post(path, headers=fixture_data["manager"], json=body).status_code == 409
    with SessionLocal() as db:
        order = db.query(PurchaseOrder).one()
        assert order.quantity == 30
        assert order.status == "Simulated Draft"
    assert (
        client.get("/api/v1/workflows", headers=fixture_data["analyst"]).json()[0]["status"]
        == "Completed"
    )


def test_export_accuracy_and_formula_protection(client, fixture_data):
    h = fixture_data["analyst"]
    report = client.post("/api/v1/reports", headers=h, json={"kind": "Daily sales MIS"}).json()
    response = client.get(f"/api/v1/reports/{report['id']}/download", headers=h)
    wb = load_workbook(io.BytesIO(response.content), data_only=True)
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)}
    assert summary["revenue"] == 5310
    assert sum(r[1] for r in list(wb["Report"].iter_rows(values_only=True))[1:-1]) == 5310
    pdf = client.get(f"/api/v1/reports/{report['id']}/download?format=pdf", headers=h)
    assert pdf.content.startswith(b"%PDF")
    assert safe_cell(' =HYPERLINK("bad")').startswith("'")
    assert safe_cell(-123) == -123
    assert (
        client.get(
            f"/api/v1/reports/{report['id']}/download", headers=fixture_data["outsider"]
        ).status_code
        == 404
    )


def test_forecast_constant_series_and_no_future_leakage(client, fixture_data):
    y = np.repeat(10, 90)
    best, forecast, baseline, scores, mae = evaluate_series(y, 30)
    assert mae == 0 and all(s["mae"] == 0 for s in scores)
    assert np.all(forecast == 10) and np.all(baseline == 10)
    # Mutating future observations cannot affect predictions from a prior origin.
    changed = y.copy()
    changed[70:] = 999
    assert np.array_equal(predict(y[:70], best, 7), predict(changed[:70], best, 7))
    with SessionLocal() as db:
        old = analytics(db, fixture_data["org"], Filters(end=AS_OF - timedelta(days=1)), {})
        assert old["metrics"]["revenue"] == 29 * 180
    result = client.get(
        "/api/v1/forecasts?start=2026-07-01&horizon=7", headers=fixture_data["analyst"]
    )
    assert result.status_code == 200, result.text
    assert len(result.json()["forecast"]) == 7


def test_workflow_failure_recovery_and_idempotency(client, fixture_data):
    h = fixture_data["analyst"]
    with SessionLocal() as db:
        row = Workflow(org_id=fixture_data["org"], kind="invalid", key="broken")
        db.add(row)
        db.commit()
        ident = row.id
    execute(ident)
    with SessionLocal() as db:
        row = db.get(Workflow, ident)
        assert row.status == "Failed" and row.attempts == 1
        row.kind = "daily-report"
        db.commit()
    assert client.post(f"/api/v1/workflows/{ident}/retry", headers=h).status_code == 200
    current = client.get("/api/v1/workflows", headers=h).json()[0]
    assert current["status"] == "Completed" and current["attempts"] == 2
    assert client.post(f"/api/v1/workflows/{ident}/retry", headers=h).status_code == 409
    body = {"kind": "daily-report", "key": "same-key"}
    first = client.post("/api/v1/workflows", headers=h, json=body).json()
    second = client.post("/api/v1/workflows", headers=h, json=body).json()
    assert first["id"] == second["id"]


def test_invalid_state_transition():
    row = Workflow(org_id="fake", kind="test", key="key", status="Completed")
    with pytest.raises(Exception):
        transition(row, "Running")


def test_restricted_analyst(client, fixture_data):
    h = fixture_data["analyst"]
    result = client.post(
        "/api/v1/analyst/query",
        headers=h,
        json={"question": "What is our current inventory value?"},
    ).json()
    assert result["rows"][0]["value"] == 5200
    assert result["mode"].startswith("Deterministic")
    assert (
        client.post(
            "/api/v1/analyst/query",
            headers=h,
            json={"question": "DROP TABLE users; place an order"},
        ).status_code
        == 422
    )
    assert (
        client.patch(
            "/api/v1/settings",
            headers=h,
            json={"safety_days": 7, "dead_days": 90, "cover_days": 14},
        ).status_code
        == 403
    )
    assert client.get("/api/v1/analytics?store_id=-1", headers=h).status_code == 422
    assert (
        client.get("/api/v1/analytics?start=2026-10-18&end=2026-01-01", headers=h).status_code
        == 422
    )
