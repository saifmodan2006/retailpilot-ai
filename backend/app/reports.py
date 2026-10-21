import io
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .analytics import analytics, inventory_metrics
from .models import Report

REPORT_TYPES = [
    "Daily sales MIS",
    "Weekly store performance",
    "Monthly management summary",
    "Category sales analysis",
    "SKU inventory",
    "Low stock",
    "Dead stock and ageing",
    "Stock cover",
    "Working capital",
]


def create_report(db, org, kind, filters, config):
    data = analytics(db, org, filters, config)
    if kind in REPORT_TYPES[4:]:
        inv = inventory_metrics(db, org, filters, config)
        if kind == "Low stock":
            inv = [r for r in inv if r["classification"] == "Low stock"]
        if kind == "Dead stock and ageing":
            inv = [r for r in inv if r["dead_stock"]]
        rows = [
            {
                k: r[k]
                for k in [
                    "sku",
                    "product",
                    "store",
                    "quantity",
                    "reserved",
                    "cost",
                    "value",
                    "daily_demand",
                    "cover_days",
                    "inactive_days",
                    "classification",
                ]
            }
            for r in inv
        ]
    else:
        rows = (
            data["trend"]
            if kind == REPORT_TYPES[0]
            else data["categories"]
            if kind == REPORT_TYPES[3]
            else data["stores"]
        )
    report = Report(
        org_id=org,
        kind=kind,
        filters={
            **filters.model_dump(mode="json"),
            "currency": config.get("currency", "INR"),
            "data_label": data["data_label"],
        },
        rows=rows,
        metrics=data["metrics"],
    )
    db.add(report)
    db.flush()
    return report


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def xlsx_export(title, rows, metrics=None, filters=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    headers = list(rows[0]) if rows else ["message"]
    ws.append(headers)
    for row in rows:
        ws.append([safe_cell(row.get(k)) for k in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="21685B")
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(46, max(16, len(header) + 4))
        for cell in list(ws.columns)[i - 1][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = (
                    "#,##0.00"
                    if header
                    in [
                        "revenue",
                        "cogs",
                        "cost",
                        "value",
                        "gross_profit",
                        "unit_price",
                    ]
                    else "0.00"
                )
        if header in ["revenue", "units", "quantity", "value", "cogs"] and rows:
            ws.cell(len(rows) + 2, i, sum(float(r.get(header) or 0) for r in rows))
            ws.cell(len(rows) + 2, i).font = Font(bold=True)
    summary = wb.create_sheet("Summary")
    summary.append(["RetailPilot AI", safe_cell(title)])
    summary.append(
        [
            "Dataset",
            safe_cell((filters or {}).get("data_label", "Fictional synthetic demonstration")),
        ]
    )
    for key, value in (filters or {}).items():
        summary.append([key, safe_cell(str(value))])
    for key, value in (metrics or {}).items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 65
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def pdf_export(report):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("RetailPilot AI | " + escape(report.kind), styles["Title"]),
        Paragraph(
            escape(report.filters.get("data_label", "Synthetic demonstration data"))
            + ". Currency: "
            + escape(report.filters.get("currency", "INR"))
            + ". "
            + escape(str(report.filters)),
            styles["Normal"],
        ),
        Spacer(1, 14),
    ]
    metrics = [
        [
            key.replace("_", " ").title(),
            f"{value:,.2f}" if isinstance(value, (int, float)) else str(value),
        ]
        for key, value in report.metrics.items()
    ]
    story.append(Table(metrics, colWidths=[220, 160]))
    story.append(Spacer(1, 18))
    if report.rows:
        headers = list(report.rows[0])
        # Full inventory exports stay legible by selecting key columns in PDF.
        headers = [
            k
            for k in headers
            if k
            in [
                "date",
                "store",
                "category",
                "sku",
                "quantity",
                "revenue",
                "units",
                "value",
                "classification",
            ]
        ][:6]
        table = Table(
            [[Paragraph(escape(k.replace("_", " ").title()), styles["Normal"]) for k in headers]]
            + [
                [Paragraph(escape(str(r.get(k, ""))), styles["Normal"]) for k in headers]
                for r in report.rows
            ],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfeee8")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(table)
    doc.build(story)
    return output.getvalue()
