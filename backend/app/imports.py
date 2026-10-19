import hashlib
import csv
import io
import re
import zipfile
from itertools import islice
from datetime import date
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import load_workbook
from fastapi import HTTPException
from sqlalchemy import select

from .models import ImportBatch, Inventory, SKU, Sale, Store, ValidationError

SALES_FIELDS = [
    "transaction_id",
    "date",
    "store",
    "sku",
    "quantity",
    "unit_price",
    "discount",
]
INVENTORY_FIELDS = ["date", "store", "sku", "quantity", "reserved", "received_date"]
ALIASES = {
    "transaction": "transaction_id",
    "transaction_no": "transaction_id",
    "invoice": "transaction_id",
    "store_code": "store",
    "sku_code": "sku",
    "qty": "quantity",
    "selling_price": "unit_price",
    "sales_date": "date",
}


def normalize(name):
    key = re.sub(r"[^a-z0-9]+", "_", str(name).lower().strip()).strip("_")
    return ALIASES.get(key, key)


def read_file(content, filename):
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(413, "Maximum upload size is 8 MB")
    try:
        if filename.lower().endswith(".csv"):
            header = next(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
            names = [normalize(c) for c in header]
            if len(names) > 50 or len(names) != len(set(names)):
                raise ValueError("Maximum 50 distinct, unambiguous columns")
            frames = {
                "CSV": pd.read_csv(
                    io.BytesIO(content), dtype=str, keep_default_na=False, nrows=20001
                )
            }
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024
                    or len(archive.infolist()) > 500
                ):
                    raise HTTPException(413, "Workbook expands beyond safe limits")
                if any("vba" in i.filename.lower() or i.flag_bits & 1 for i in archive.infolist()):
                    raise ValueError("Macros and encrypted workbooks are not supported")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            count = 0
            try:
                for sheet in workbook:
                    if sheet.max_column and sheet.max_column > 50:
                        raise ValueError("Maximum 50 columns per sheet")
                    for i, cells in enumerate(islice(sheet.iter_rows(), 20002)):
                        if i == 0:
                            names = [normalize(c.value) for c in cells if c.value is not None]
                            if len(names) != len(set(names)):
                                raise ValueError("Duplicate or ambiguous column names")
                        else:
                            count += 1
                        if count > 20000:
                            raise HTTPException(413, "Maximum 20,000 rows per upload")
                        if any(c.data_type == "f" for c in cells):
                            raise ValueError(
                                "Formula cells are not accepted. Upload a values-only workbook."
                            )
            finally:
                workbook.close()
            frames = pd.read_excel(
                io.BytesIO(content),
                sheet_name=None,
                dtype=str,
                keep_default_na=False,
                engine="openpyxl",
            )
        else:
            raise HTTPException(415, "Only .csv and .xlsx are supported")
        if sum(len(f) for f in frames.values()) > 20000:
            raise HTTPException(413, "Maximum 20,000 rows per upload")
        raw = []
        columns = set()
        for sheet, frame in frames.items():
            if len(set(frame.columns)) != len(frame.columns):
                raise ValueError("Duplicate column names")
            columns.update(str(c) for c in frame.columns)
            for i, row in enumerate(frame.to_dict("records"), 2):
                raw.append(
                    {
                        "sheet": sheet,
                        "row": i,
                        "values": {str(k): str(v).strip() for k, v in row.items()},
                    }
                )
        if not raw:
            raise ValueError("File has no data rows")
        return raw, sorted(columns)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"Cannot parse spreadsheet: {str(exc)[:150]}") from exc


def number(value, default=None):
    if value == "" and default is not None:
        return Decimal(str(default))
    try:
        n = Decimal(str(value).replace("₹", "").replace("INR", "").replace(",", "").strip())
        if not n.is_finite() or abs(n) > Decimal("10000000"):
            raise ValueError("Numeric value outside supported range")
        return n
    except InvalidOperation as exc:
        raise ValueError("Invalid numeric value") from exc


def parse_date(value):
    value = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}( 00:00:00)?", value):
        return date.fromisoformat(value[:10])
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", value):
        return pd.to_datetime(value, dayfirst=True, errors="raise").date()
    raise ValueError("Date must be YYYY-MM-DD or DD/MM/YYYY without a time or suffix")


def validate(db, org, raw, mapping, kind):
    stores = {s.code: s for s in db.scalars(select(Store).where(Store.org_id == org))}
    skus = {s.code: s for s in db.scalars(select(SKU).where(SKU.org_id == org))}
    required = ["date", "store", "sku", "quantity"] + (
        ["transaction_id", "unit_price"] if kind == "sales" else []
    )
    missing = set(required) - set(mapping.values())
    if missing:
        return [], [
            {
                "row": 1,
                "sheet": "Headers",
                "message": "Missing required columns: " + ", ".join(sorted(missing)),
            }
        ]
    existing = (
        set(db.scalars(select(Sale.transaction_id).where(Sale.org_id == org)))
        if kind == "sales"
        else set()
    )
    seen = set()
    valid, errors = [], []
    for source in raw:
        row = {
            target: source["values"].get(column, "") for column, target in mapping.items() if target
        }
        try:
            if any(row.get(k, "") == "" for k in required):
                raise ValueError("Missing required value")
            if row["store"] not in stores:
                raise ValueError("Unknown store code")
            if row["sku"] not in skus:
                raise ValueError("Unknown SKU")
            day = parse_date(row["date"])
            if day > date.today():
                raise ValueError("Future business date is not allowed")
            qty = number(row["quantity"])
            if qty != int(qty) or abs(qty) > 1000000:
                raise ValueError("Quantity must be an integer")
            out = {
                "date": day.isoformat(),
                "store_id": stores[row["store"]].id,
                "sku_id": skus[row["sku"]].id,
                "quantity": int(qty),
            }
            if kind == "sales":
                key = row["transaction_id"]
                if len(key) > 120 or key in existing or key in seen:
                    raise ValueError("Duplicate or invalid transaction ID")
                price, discount = (
                    number(row["unit_price"]),
                    number(row.get("discount", ""), 0),
                )
                if price < 0 or not 0 <= discount <= 1 or qty == 0:
                    raise ValueError(
                        "Price must be nonnegative; discount is 0–1; quantity cannot be zero"
                    )
                if abs(qty * price * (1 - discount)) > Decimal("999999999999.99") or abs(
                    qty * skus[row["sku"]].cost
                ) > Decimal("999999999999.99"):
                    raise ValueError("Transaction amount exceeds the supported monetary range")
                out.update(
                    transaction_id=key,
                    revenue=float((qty * price * (1 - discount)).quantize(Decimal(".01"))),
                    cogs=float(qty * skus[row["sku"]].cost),
                )
            else:
                reserved = number(row.get("reserved", ""), 0)
                received = parse_date(row.get("received_date") or row["date"])
                if (
                    qty < 0
                    or reserved < 0
                    or reserved != int(reserved)
                    or reserved > qty
                    or received > day
                ):
                    raise ValueError("Invalid stock, reserved quantity, or received date")
                key = (out["store_id"], out["sku_id"], out["date"])
                if key in seen:
                    raise ValueError("Duplicate inventory snapshot")
                out.update(reserved=int(reserved), received_date=received.isoformat())
            seen.add(key)
            valid.append(out)
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append(
                {
                    "row": source["row"],
                    "sheet": source["sheet"],
                    "message": str(exc)[:240],
                }
            )
    return valid, errors


def preview(db, org, content, filename, kind):
    raw, columns = read_file(content, filename)
    digest = hashlib.sha256(content + kind.encode()).hexdigest()
    batch = db.scalar(
        select(ImportBatch).where(ImportBatch.org_id == org, ImportBatch.fingerprint == digest)
    )
    if batch:
        return batch
    allowed = SALES_FIELDS if kind == "sales" else INVENTORY_FIELDS
    mapping = {c: normalize(c) for c in columns if normalize(c) in allowed}
    if len(mapping.values()) != len(set(mapping.values())):
        # Keep source columns available but require the user to disambiguate aliases.
        used = set()
        for column in list(mapping):
            if mapping[column] in used:
                mapping[column] = ""
            else:
                used.add(mapping[column])
    batch = ImportBatch(
        org_id=org,
        fingerprint=digest,
        filename=filename[:200],
        kind=kind,
        payload=raw,
        mapping=mapping,
    )
    db.add(batch)
    db.flush()
    revalidate(db, batch, mapping)
    return batch


def revalidate(db, batch, mapping):
    if batch.status == "Imported":
        raise HTTPException(409, "Imported batches cannot be remapped")
    valid, errors = validate(db, batch.org_id, batch.payload, mapping, batch.kind)
    db.query(ValidationError).filter(
        ValidationError.batch_id == batch.id, ValidationError.org_id == batch.org_id
    ).delete()
    db.add_all([ValidationError(org_id=batch.org_id, batch_id=batch.id, **e) for e in errors])
    batch.mapping = mapping
    batch.accepted, batch.rejected = len(valid), len(errors)
    db.flush()
    return valid


def commit_import(db, batch):
    if batch.status == "Imported":
        return False
    valid = revalidate(db, batch, batch.mapping)
    if not valid:
        raise HTTPException(422, "No accepted rows. Review validation errors and column mappings.")
    for row in valid:
        row["date"] = date.fromisoformat(row["date"])
        if batch.kind == "sales":
            db.add(Sale(org_id=batch.org_id, **row))
        else:
            row["received_date"] = date.fromisoformat(row["received_date"])
            previous = db.scalar(
                select(Inventory).where(
                    Inventory.org_id == batch.org_id,
                    Inventory.store_id == row["store_id"],
                    Inventory.sku_id == row["sku_id"],
                    Inventory.date == row["date"],
                )
            )
            if previous:
                previous.quantity, previous.reserved, previous.received_date = (
                    row["quantity"],
                    row["reserved"],
                    row["received_date"],
                )
            else:
                db.add(Inventory(org_id=batch.org_id, **row))
    batch.status = "Imported"
    db.flush()
    return True
